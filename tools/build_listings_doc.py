#!/usr/bin/env python3
"""Build SAP_BDC_360_Marketplace_Listings.docx from /tmp/listings.json.

    python3 tools/collect_listings.py && python3 tools/build_listings_doc.py
"""
import datetime as dt
import json
import pathlib
import sys

from docx import Document
from docx.shared import Pt

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from build_asset_index_links import link_cell  # noqa: E402
from docx_kit import GREY, SAP_NAVY, SNOW_BLUE, body, bullet, h1, h2, setup_page, table  # noqa: E402

OUT = pathlib.Path.home() / "Downloads" / "SAP_BDC_360_Marketplace_Listings.docx"
L = json.loads(pathlib.Path("/tmp/listings.json").read_text())
REG = {"dfreriksdemo": ("US West (Oregon)", "AWS us-west-2"),
       "dfreriks_eu_demo": ("EU (Frankfurt)", "AWS eu-central-1"),
       "dfreriks_apac_demo": ("APAC (Sydney)", "AWS ap-southeast-2")}
ORDER = list(REG)


def acct_url(n):
    return f"https://app.snowflake.com/{L[n]['org'].lower()}/{L[n]['account'].lower()}/"


def product(title):
    return title.replace("SAP BDC ", "SAP ").replace(" Demo", "")


idx = {n: {r["name"]: r for r in L[n]["rows"] if "360" in r["title"] and r["state"] == "PUBLISHED"} for n in ORDER}
groups = {}
for n in ORDER:
    for name, r in idx[n].items():
        groups.setdefault(product(r["title"]), set()).add(name)

doc = Document()
setup_page(doc)
for t, sz, b, c, a in (("SAP BDC 360 Demo Assets", 20, True, SAP_NAVY, 2),
                       ("Internal Marketplace listings — all three regions", 11.5, False, SNOW_BLUE, 2),
                       (f"Read live from each account on {dt.date.today():%d %B %Y}  ·  Owner: Dave Freriks", 9, False, GREY, 14)):
    p = doc.add_paragraph(); r = p.add_run(t); r.font.size = Pt(sz); r.font.bold = b; r.font.color.rgb = c
    p.paragraph_format.space_after = Pt(a)
body(doc, "Every SAP BDC 360 demo is published to the Snowflake Internal Marketplace in three regions, as two "
          "listings: the Native App (the full dashboard with its data bundled, nothing to set up) and a data "
          "share (the analytics tables, semantic view and Cortex Agent, to query from your own account or use "
          "in Snowflake Intelligence). Open the Internal Marketplace in the account for your region and search "
          "for the listing title or global name below.")
h1(doc, "Regional accounts")
t = table(doc, ["Region", "Cloud region", "Account (Snowsight)"],
          [[REG[n][0], REG[n][1], acct_url(n)] for n in ORDER], widths=[1.5, 1.5, 3.7])
for i, n in enumerate(ORDER, start=1):  # make the account addresses clickable
    link_cell(t.rows[i].cells[2], acct_url(n))

h1(doc, "Listings by product")
missing = []
for prod in sorted(groups):
    h2(doc, prod)
    names = sorted(groups[prod], key=lambda x: 0 if any(idx[n].get(x, {}).get("is_application") == "true" for n in ORDER) else 1)
    rows = []
    for name in names:
        ref = next(idx[n][name] for n in ORDER if name in idx[n])
        kind = "Native App" if ref["is_application"] == "true" else "Data share"
        for n in ORDER:
            r = idx[n].get(name)
            if not r:
                missing.append(f"{prod} {kind.lower()}: {REG[n][0]}")
            rows.append([kind if n == ORDER[0] else "", REG[n][0], r["global_name"] if r else "not published",
                         r["uniform_listing_locator"] if r and r["uniform_listing_locator"] != "None" else "—"])
    table(doc, ["Listing", "Region", "Global name", "Uniform listing locator"], rows, widths=[1.0, 1.4, 1.3, 3.0])

h1(doc, "Notes")
for m in missing:
    bullet(doc, f"Not published: {m}.")
bullet(doc, "Global names differ by region because each region's listing is a separate object. The uniform "
            "listing locator (ORGDATACLOUD$INTERNAL$…) is the same in every region.")
bullet(doc, "Spend, People and Working Capital data shares in EU and APAC carry a copy of the US analytics layer "
            "(their SAP BDC source shares exist only in the US account). The copies do not auto-refresh; "
            "re-run sap-bdc-spend-360/tools/publish_data_share.py after the US data changes.")
bullet(doc, "Draft and test listings in the US account are left out.")
doc.save(OUT)
print(OUT, sum(len(v) for v in groups.values()), "listings;", len(missing), "not published")
