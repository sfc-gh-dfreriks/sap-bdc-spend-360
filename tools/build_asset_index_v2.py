#!/usr/bin/env python3
"""Build the SAP BDC 360 demo asset index, V2: the V1 index plus dashboard
screenshots, headline KPIs, page inventories, use cases and the data/AI stack.

Inputs, all read live beforehand:
  /tmp/listings.json, /tmp/apps.json   tools/collect_listings.py (the three accounts)
  /tmp/portfolio/portfolio.json + PNGs tools/capture_portfolio.py (the running demos)
Repo facts, use cases and kit files are read from disk here.
"""
import datetime as dt
import json
import pathlib
import re
import subprocess
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

sys.path.insert(0, str(pathlib.Path(__file__).parent))
from build_asset_index_links import link_cell  # noqa: E402
from docx_kit import (GREY, LIGHT_HEX, NAVY_HEX, SAP_NAVY, SNOW_BLUE, body, bullet,  # noqa: E402
                      fixed, h1, h2, setup_page, shade, table)

HOME = pathlib.Path.home()
SK = HOME / "Documents/SAP/SAP Skills"
SAP = HOME / "Documents/SAP"
GH = "https://github.com/sfc-gh-dfreriks/"
PAGES = "https://sfc-gh-dfreriks.github.io/"
OUT = HOME / "Downloads/SAP_BDC_360_Demo_Asset_Index_V2.docx"
SHOTS = pathlib.Path("/tmp/portfolio")
IMG = SHOTS / "_docx"
W = 6.9  # usable page width, inches

L = json.loads(pathlib.Path("/tmp/listings.json").read_text())
A = json.loads(pathlib.Path("/tmp/apps.json").read_text())
PF = json.loads((SHOTS / "portfolio.json").read_text())
REG = [("dfreriksdemo", "US West (Oregon)"), ("dfreriks_eu_demo", "EU (Frankfurt)"),
       ("dfreriks_apac_demo", "APAC (Sydney)")]

# key, title, blurb, asset repo, public repo, local app dir, ports, app, listings, kit dir
P = [
    ("finance", "SAP Finance 360", "GL, cost and profit centres, AP/AR and period analysis on SAP finance data.",
     "sap-bdc-finance-360", "finance-360-public", "finance_dashboard_react", "3002 / 5175",
     "FINANCE_360_APP", ["FINANCE_360_ORG", "SAP_BDC_FINANCE_360_DEMO"], "Finance_360_Presales_Kit"),
    ("sales", "SAP Sales 360", "Pipeline, funnel, customer health, products, leaderboard and forecast.",
     "sap-bdc-sales-360", "sales-360-public", "sales_360_react", "3003 / 5176",
     "SALES_360_APP", ["SALES_360_ORG", "SAP_BDC_SALES_360_DEMO"], "Sales_360_Presales_Kit"),
    ("supply", "SAP Supply Chain 360", "Production orders, inventory, deliveries, supplier quality and work centres.",
     "sap-bdc-supply-chain-360", "supply-chain-360-public", "../SAP-Qlikstarts/supply_chain_dashboard_react",
     "3001 / 5174", "SUPPLY_CHAIN_360_APP", ["SUPPLY_CHAIN_360_ORG", "SAP_BDC_SUPPLY_CHAIN_360"],
     "Supply_Chain_360_Presales_Kit"),
    ("ontology", "SAP Supply Chain Ontology", "Ontology layer over supply chain data: classes, relations, knowledge graph and disruption scenarios.",
     "supply-chain-ontology", "supply-chain-ontology", "supply_chain_ontology_react", "3009 / 5179",
     None, [], "Supply_Chain_Ontology_Presales_Kit"),
    ("people", "SAP People 360", "Headcount, attrition, compensation, diversity, performance, learning and recruiting.",
     "sap-bdc-people-360", "people-360-public", "people_360_react", "3006 / 5180",
     "PEOPLE_360_APP", ["PEOPLE_360_ORG", "SAP_BDC_PEOPLE_360_DEMO"], "People_360_Presales_Kit"),
    ("spend", "SAP Spend 360", "Spend by category and supplier, contract compliance, supplier risk, ESG and savings, with currency conversion.",
     "sap-bdc-spend-360", "spend-360-public", "spend_360_react", "3007 / 5181",
     "SPEND_360_APP", ["SPEND_360_ORG", "SAP_BDC_SPEND_360_DEMO"], "Spend_360_Presales_Kit"),
    ("wc", "SAP Working Capital 360", "Receivables, payables and inventory working capital: DSO, DPO, DIO and cash conversion.",
     "sap-bdc-working-capital-360", "working-capital-360-public", "working_capital_360_react", "3010 / 5185",
     "WORKING_CAPITAL_360_APP", ["WORKING_CAPITAL_360_ORG", "SAP_BDC_WORKING_CAPITAL_360_DEMO"], "Working_Capital_360_Presales_Kit"),
]

# Band colour per demo, so each section is recognisable at a glance.
ACCENT = {"finance": "1B3A57", "sales": "0E7C86", "supply": "2E5AAC", "ontology": "5B3F8C",
          "people": "B4580E", "spend": "1B7F4B", "wc": "9A2A4A", "erp": "3C4A5A"}

# Data and AI stack. The SAP BDC data products are summarised from each repo's
# sql/01_l0_sources.md; semantic views and agents come from SHOW AGENTS / USE_CASES.md.
STACK = {
    "finance": ("SAP S/4HANA Journal Entry header and items, Supplier Invoice, GL Account, Cost Center, Profit Center",
                "SAP_FINANCE_360.ANALYTICS.SAP_FINANCE_360", "SAP_FINANCE_360.ANALYTICS.SAP_FINANCE_360_AGENT"),
    "sales": ("SAP S/4HANA Sales Order and item, Customer, Product; CRM Opportunity, Activity and Sales Rep",
              "SAP_SALES_360.SEMANTIC.SAP_SALES_360_ANALYTICS", "SAP_SALES_360.SEMANTIC.SAP_SALES_360_AGENT"),
    "supply": ("SAP S/4HANA Plant, Production Order confirmation and version, BOM, Material Stock, Delivery, "
               "Shipping Point, Storage Location, Project/WBS, Supplier Quality, supply-chain nodes and flows",
               "SAP_SUPPLY_CHAIN.ANALYTICS.SAP_SUPPLY_CHAIN_360", "SAP_SUPPLY_CHAIN.ANALYTICS.SAP_SC360_ANALYST_AGENT"),
    "ontology": ("The SAP BDC supply-chain catalog slice (36 data products, 338 CDS entities) mapped to 15 ontology classes "
                 "over the Supply Chain 360 data",
                 "3 semantic views in SAP_SUPPLY_CHAIN.ONTOLOGY", "SAP_SUPPLY_CHAIN.ONTOLOGY.SUPPLY_CHAIN_AGENT"),
    "people": ("SAP SuccessFactors Core Workforce standard data product (COREWORKFORCE_STANDARDFIELDS)",
               "SAP_PEOPLE_360.SEMANTIC.SAP_PEOPLE_360_ANALYTICS", "SAP_PEOPLE_360.AGENTS.SAP_PEOPLE_ANALYST"),
    "spend": ("SAP S/4HANA Purchase Order item and Supplier, plus Spend Intelligence enrichment (risk, ESG, "
              "category taxonomy, savings) and ECB daily FX rates",
              "SAP_SPEND_360.SEMANTIC.SAP_SPEND_360_ANALYTICS", "SAP_SPEND_360.AGENTS.SAP_SPEND_ANALYST"),
    "wc": ("SAP S/4HANA Entry View Journal Entry, Supplier Invoice, Payment Terms, Billing Document, Collections "
           "(worklist, dispute, promise to pay), Cash Flow, Bank Account, Physical Inventory Document",
           "SAP_WORKING_CAPITAL_360.SEMANTIC.SAP_WORKING_CAPITAL_360_ANALYTICS",
           "SAP_WORKING_CAPITAL_360.AGENTS.SAP_WORKING_CAPITAL_ANALYST"),
    "erp": ("No sources of its own: joins the Finance 360, Spend 360 and People 360 analytics layers by company",
            "SAP_ERP_360.SEMANTIC.SAP_ERP_360_ANALYTICS", "SAP_ERP_360.AGENTS.SAP_ERP_ANALYST"),
}


# --------------------------------------------------------------- docx helpers
def hyperlink_run(paragraph, url, text, size=8.5):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                         is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "0563C1"), ("w:u", "single"), ("w:sz", str(int(size * 2)))):
        e = OxmlElement(tag); e.set(qn("w:val"), val); rpr.append(e)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); paragraph._p.append(h)


def link_table(doc, headers, rows, widths):
    """rows: list of cells; a cell may be ("url", link, text) to render as a hyperlink."""
    plain = [[c[2] if isinstance(c, tuple) else c for c in row] for row in rows]
    t = table(doc, headers, plain, widths=widths, size=8.5, zebra=True)
    for i, row in enumerate(rows, start=1):
        for j, c in enumerate(row):
            if isinstance(c, tuple) and c[1]:
                link_cell(t.rows[i].cells[j], c[1], c[2])
    return t


def no_borders(t):
    tbl_pr = t._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "nil"); b.append(e)
    tbl_pr.append(b)


def cell_margins(cell, pts=5):
    tc_pr = cell._tc.get_or_add_tcPr()
    m = OxmlElement("w:tcMar")
    for side in ("top", "bottom", "left", "right"):
        e = OxmlElement(f"w:{side}"); e.set(qn("w:w"), str(pts * 20)); e.set(qn("w:type"), "dxa"); m.append(e)
    tc_pr.append(m)


def run(p, text, size, bold=False, color=None, italic=False):
    r = p.add_run(text); r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    if color is not None:
        r.font.color.rgb = color
    return r


def band(doc, title, sub, hexcolor):
    """Full-width coloured header band for a demo section."""
    t = doc.add_table(rows=1, cols=1); no_borders(t)
    c = t.rows[0].cells[0]; shade(c, hexcolor); cell_margins(c, 8)
    p = c.paragraphs[0]; run(p, title, 17, True, RGBColor(0xFF, 0xFF, 0xFF))
    q = c.add_paragraph(); run(q, sub, 9.5, False, RGBColor(0xE6, 0xEE, 0xF4))
    fixed(t, [W])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def kpi_tiles(doc, kpis, hexcolor, cols=4):
    """KPI cards: big value, small label, one-line note; rows of `cols`."""
    if not kpis:
        return
    rows = (len(kpis) + cols - 1) // cols
    t = doc.add_table(rows=rows, cols=cols); no_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(rows * cols):
        c = t.rows[i // cols].cells[i % cols]; cell_margins(c, 5)
        if i >= len(kpis):
            continue
        k = kpis[i]; shade(c, LIGHT_HEX)
        p = c.paragraphs[0]; run(p, k["label"].upper(), 7, True, GREY)
        p.paragraph_format.space_after = Pt(1)
        q = c.add_paragraph(); run(q, k["value"], 16, True, RGBColor.from_string(hexcolor))
        q.paragraph_format.space_after = Pt(0)
        if k.get("note"):
            n = c.add_paragraph(); run(n, k["note"], 7, False, GREY)
    fixed(t, [W / cols] * cols)
    # white gutter between tiles
    for row in t.rows:
        for c in row.cells:
            tc_pr = c._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
            for side in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{side}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "24")
                e.set(qn("w:color"), "FFFFFF"); b.append(e)
            tc_pr.append(b)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def docx_img(src, width_px=1700):
    """Downscale a PNG capture to a compact JPEG for the document."""
    IMG.mkdir(exist_ok=True)
    out = IMG / (pathlib.Path(src).parent.name + "_" + pathlib.Path(src).stem + ".jpg")
    im = Image.open(src).convert("RGB")
    if im.width > width_px:
        im = im.resize((width_px, round(im.height * width_px / im.width)), Image.LANCZOS)
    im.save(out, quality=82, optimize=True)
    return str(out)


def framed_picture(cell_or_doc, path, width, caption=None):
    p = cell_or_doc.add_paragraph() if hasattr(cell_or_doc, "add_paragraph") else cell_or_doc
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    # thin grey frame around the picture
    for pic in p._p.iter(qn("pic:spPr")):
        ln = OxmlElement("a:ln"); ln.set("w", "9525")
        sf = OxmlElement("a:solidFill"); clr = OxmlElement("a:srgbClr"); clr.set("val", "B8C4CE")
        sf.append(clr); ln.append(sf); pic.append(ln)
    p.paragraph_format.space_after = Pt(2)
    if caption:
        q = cell_or_doc.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(q, caption, 8, False, GREY, italic=True); q.paragraph_format.space_after = Pt(6)


def toc(doc):
    p = doc.add_paragraph(); r = p.add_run()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None), (None, None), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r._r.append(f)
        elif text:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text; r._r.append(it)
        else:
            t = OxmlElement("w:t"); t.text = "Right-click and choose Update Field to build the table of contents."; r._r.append(t)
    uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); doc.settings.element.append(uf)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def footer(doc):
    p = doc.sections[0].footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "SAP BDC 360 Demo Portfolio · Asset Index V2 · page ", 8, False, GREY)
    r = p.add_run(); r.font.size = Pt(8)
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r._r.append(f)
        else:
            it = OxmlElement("w:instrText"); it.text = text; r._r.append(it)


# ------------------------------------------------------------------ facts
def repo_meta(name):
    if not name:
        return None
    for base in (SK, HOME / "Documents/GitHub"):
        d = base / name
        if (d / ".git").exists():
            g = lambda *a: subprocess.run(["git", "-C", str(d), *a], capture_output=True, text=True).stdout.strip()  # noqa: E731
            return {"path": str(d), "branch": g("branch", "--show-current"), "last": g("log", "-1", "--format=%cs")}
    return {"path": "(not cloned locally)", "branch": "", "last": ""}


VIS = {}
try:
    out = subprocess.run(["gh", "repo", "list", "sfc-gh-dfreriks", "--limit", "200", "--json", "name,visibility"],
                         capture_output=True, text=True, timeout=60).stdout
    VIS = {r["name"]: r["visibility"].title() for r in json.loads(out or "[]")}
except Exception:  # noqa: BLE001
    pass


def listing(conn, name):
    return next((r for r in L[conn]["rows"] if r["name"] == name and r["state"] == "PUBLISHED"), None)


def app(conn, name):
    return next((a for a in A.get(conn, []) if a["app"] == name), None)


def kit_files(kit):
    d = SAP / kit if kit else None
    if not d or not d.exists():
        return []
    return sorted(str(p.relative_to(d)) for p in d.rglob("*")
                  if p.suffix.lower() in (".docx", ".pptx", ".mp4", ".pdf", ".xlsx"))


def acct_url(conn):
    return f"https://app.snowflake.com/{L[conn]['org'].lower()}/{L[conn]['account'].lower()}/"


def use_cases(key, repo):
    """Parse docs/USE_CASES.md: the summary table plus each section's business question."""
    path = SK / (repo or "erp_360_react") / "docs/USE_CASES.md"
    if key == "erp":
        path = SK / "erp_360_react/docs/USE_CASES.md"
    if not path.exists():
        return []
    text = path.read_text()
    questions = {}
    for m in re.finditer(r"^## (\d+)\. .*?$(.*?)(?=^## |\Z)", text, re.M | re.S):
        q = re.search(r"\*\*Business question:\*\*\s*(.+)", m.group(2))
        questions[m.group(1)] = q.group(1).strip() if q else ""
    rows = []
    for m in re.finditer(r"^\| (\d+) \| ([^|]+)\| ([^|]+)\| ([^|]+)\|", text, re.M):
        n = m.group(1)
        rows.append([n, m.group(2).strip(), m.group(3).strip(), questions.get(n, "")])
    return rows


def ncount(pattern):
    return sum(1 for conn, _ in REG for r in L[conn]["rows"] if r["state"] == "PUBLISHED" and re.search(pattern, r["name"]))


# ------------------------------------------------------------------ build
doc = Document()
setup_page(doc)
footer(doc)
all_uc = {k: use_cases(k, repo) for k, _, _, repo, *_ in P}
n_apps = sum(1 for k, *rest in P if rest[6] for conn, _ in REG if app(conn, rest[6]))
n_list = sum(1 for k, *rest in P for n in rest[7] for conn, _ in REG if listing(conn, n))
n_pages = sum(len(PF.get(k, {}).get("nav", [])) for k, *_ in P)

# cover
cov = doc.add_table(rows=1, cols=1); no_borders(cov)
c = cov.rows[0].cells[0]; shade(c, NAVY_HEX); cell_margins(c, 14)
run(c.paragraphs[0], "SAP BDC 360 Demo Portfolio", 26, True, RGBColor(0xFF, 0xFF, 0xFF))
run(c.add_paragraph(), "Asset Index · V2", 14, False, SNOW_BLUE)
run(c.add_paragraph(), "Every demo, what is inside it, what it shows, and where to find it: repositories, public sites, "
    "Native Apps, Marketplace listings and presales kits", 10, False, RGBColor(0xE6, 0xEE, 0xF4))
fixed(cov, [W])
doc.add_paragraph()

# headline numbers
kpi_tiles(doc, [
    {"label": "Demos", "value": str(len(P)), "note": "6 lines of business + ontology"},
    {"label": "Dashboard pages", "value": str(n_pages), "note": "across the seven apps"},
    {"label": "Documented use cases", "value": str(sum(len(v) for v in all_uc.values())), "note": "persona + business question"},
    {"label": "Regions", "value": "3", "note": "US West · EU · APAC"},
    {"label": "Native App installs", "value": str(n_apps), "note": "6 apps × 3 regional accounts"},
    {"label": "Marketplace listings", "value": str(n_list), "note": "apps + data shares, published"},
    {"label": "Cortex agents", "value": str(len(P)), "note": "one per demo"},
    {"label": "Public demo sites", "value": str(sum(1 for p in P if p[4])), "note": "GitHub Pages, no login"},
], NAVY_HEX)

# cover mosaic of overview screenshots, 4 x 2
mos = doc.add_table(rows=(len(P) + 3) // 4, cols=4); no_borders(mos)
for i, (k, title, *_) in enumerate(P):
    cell = mos.rows[i // 4].cells[i % 4]; cell_margins(cell, 3)
    shot = PF.get(k, {}).get("shots", [None])[0]
    if shot:
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.paragraphs[0].add_run().add_picture(docx_img(shot, 900), width=Inches(W / 4 - 0.12))
        q = cell.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(q, title.replace("SAP ", ""), 7.5, True, SAP_NAVY)
fixed(mos, [W / 4] * 4)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, f"Read live from GitHub, the running demos and the three demo accounts on {dt.date.today():%d %B %Y}  ·  Owner: Dave Freriks",
    8.5, False, GREY)
page_break(doc)

h1(doc, "Contents")
toc(doc)
body(doc, "Section index", italic=True)
for line in (["1. Portfolio at a glance", "2. Headline KPIs", "3. Demo accounts by region", "4. Demos"]
             + [f"    4.{i} {p[1]}" for i, p in enumerate(P, start=1)]
             + ["5. Shared tooling, skills and cross-demo documents", "6. Known gaps"]):
    q = doc.add_paragraph(line); q.paragraph_format.space_after = Pt(1)
page_break(doc)

# 1. coverage matrix
h1(doc, "1. Portfolio at a glance")
body(doc, "One row per demo. ✓ means the asset exists and was verified today; — means it does not exist yet.")
rows = []
for k, title, _, repo, pub, local, ports, appn, lst, kit in P:
    regions = [lbl.split(" ")[0] for conn, lbl in REG if appn and app(conn, appn)]
    rows.append([title, str(len(PF.get(k, {}).get("nav", []))), str(len(all_uc[k]) or "—"),
                 "✓" if repo else "—", "✓" if pub else "—", ", ".join(regions) or "—",
                 "✓" if any(listing(c, n) for c, _ in REG for n in lst) else "—", "✓" if kit_files(kit) else "—"])
table(doc, ["Demo", "Pages", "Use cases", "Repo", "Public site", "Native App", "Listings", "Kit"], rows,
      widths=[1.7, 0.55, 0.7, 0.55, 0.75, 1.15, 0.7, 0.5], zebra=True)

body(doc, "How every demo is built", size=10.5, color=SAP_NAVY)
table(doc, ["Layer", "What it is", "Where"], [
    ["L0 · SAP BDC", "SAP standard data products shared zero-copy through BDC Connect", "SAP_BDC_DEMO_* catalog-linked databases"],
    ["L1 · Curated", "Passthrough views that rename and type the SAP fields", "<DEMO>.SAP_BDC_L1"],
    ["L2 · Analytics", "Dynamic tables that shape KPIs, joins and enrichment", "<DEMO>.ANALYTICS"],
    ["AI", "Semantic view for Cortex Analyst and a Cortex Agent per demo", "<DEMO>.SEMANTIC / .AGENTS"],
    ["Apps", "React + Node app, run locally, as a public snapshot and as a Native App on SPCS", "GitHub, Pages, 3 accounts"],
], widths=[1.3, 3.5, 2.1], zebra=True)

# 2. KPI summary
page_break(doc)
h1(doc, "2. Headline KPIs")
body(doc, "The first KPI cards on each demo's overview page, read from the running app today. The figures are demo data "
     "shaped to look like a mid-size manufacturer; they are consistent across demos that share sources.")
rows = []
for k, title, *_ in P:
    ks = PF.get(k, {}).get("kpis", [])[:4]
    rows.append([title] + [f"**{x['value']}**  {x['label']}" for x in ks] + [""] * (4 - len(ks)))
table(doc, ["Demo", "KPI 1", "KPI 2", "KPI 3", "KPI 4"], rows, widths=[1.5, 1.35, 1.35, 1.35, 1.35], size=8.5, zebra=True)

# 3. accounts
h1(doc, "3. Demo accounts by region")
link_table(doc, ["Region", "Connection", "Organisation / account", "Snowsight"],
           [[lbl, conn, f"{L[conn]['org']} / {L[conn]['account']}", ("url", acct_url(conn), acct_url(conn))] for conn, lbl in REG],
           widths=[1.3, 1.3, 2.0, 2.3])

# 4. per demo
page_break(doc)
h1(doc, "4. Demos")
for i, (k, title, blurb, repo, pub, local, ports, appn, lst, kit) in enumerate(P, start=1):
    if i > 1:
        page_break(doc)
    pf = PF.get(k, {})
    color = ACCENT[k]
    h2(doc, f"4.{i} {title}")
    band(doc, title, blurb, color)

    kpi_tiles(doc, pf.get("kpis", [])[:8], color)
    shots = pf.get("shots", [])
    labels = pf.get("shot_labels", [])
    if shots:
        framed_picture(doc, docx_img(shots[0]), 5.9, f"{title} · {labels[0] if labels else 'Overview'}")

    # what's inside
    sv, ag = STACK[k][1], STACK[k][2]
    wi = body(doc, "What's inside", size=11, color=RGBColor.from_string(color), after=2)
    nav = pf.get("nav", [])
    half = (len(nav) + 1) // 2
    t = doc.add_table(rows=1, cols=2); no_borders(t)
    for col, items in enumerate((nav[:half], nav[half:])):
        cell = t.rows[0].cells[col]
        for j, name in enumerate(items):
            q = cell.paragraphs[0] if j == 0 else cell.add_paragraph()
            run(q, f"{nav.index(name) + 1:>2}.  ", 9, True, RGBColor.from_string(color)); run(q, name, 9)
            q.paragraph_format.space_after = Pt(1)
    fixed(t, [W / 2, W / 2])
    for cell in t.rows[0].cells:
        for q in cell.paragraphs:
            q.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # feature screenshots, side by side
    if len(shots) > 1:
        g = doc.add_table(rows=1, cols=2); no_borders(g)
        for j, s in enumerate(shots[1:3]):
            cell = g.rows[0].cells[j]; cell_margins(cell, 3)
            framed_picture(cell.paragraphs[0], docx_img(s, 1200), W / 2 - 0.15)
            q = cell.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run(q, labels[j + 1] if j + 1 < len(labels) else "", 8, False, GREY, italic=True)
        fixed(g, [W / 2, W / 2])
        doc.add_paragraph().paragraph_format.space_after = Pt(0)

    body(doc, "Data and AI", size=11, color=RGBColor.from_string(color), after=2)
    table(doc, ["Layer", "Objects"], [
        ["SAP BDC sources", STACK[k][0]],
        ["Semantic view", f"`{sv}`" if "." in sv else sv],
        ["Cortex agent", f"`{ag}`"],
    ], widths=[1.4, 5.5], size=8.5)

    uc = all_uc[k]
    if uc:
        body(doc, "Use cases", size=11, color=RGBColor.from_string(color), after=2)
        table(doc, ["#", "Use case", "Persona", "Business question"], uc, widths=[0.3, 1.6, 1.6, 3.4], size=8.5, zebra=True)

    body(doc, "Links and locations", size=11, color=RGBColor.from_string(color), after=2)
    links = []
    rm = repo_meta(repo)
    if repo:
        links.append(["Source repository", ("url", GH + repo, GH + repo),
                      f"{VIS.get(repo, '?')} · {rm['branch'] or '—'} · {rm['last'] or '—'}"])
    if pub:
        links.append(["Public demo", ("url", PAGES + pub + "/", PAGES + pub + "/"), "Static snapshot, no login"])
        if pub != repo:
            links.append(["Public site repository", ("url", GH + pub, GH + pub), VIS.get(pub, "?")])
    links.append(["Local app", str((SK / local).resolve()), f"ports {ports}"])
    link_table(doc, ["Asset", "Link / location", "Notes"], links, widths=[1.4, 3.8, 1.7])

    if appn:
        rws = []
        for conn, lbl in REG:
            a = app(conn, appn)
            rws.append([lbl, ("url", "https://" + a["url"], a["url"]) if a and a["url"] else "not installed",
                        f"{a['version']} patch {a['patch']}" if a else "—"])
        body(doc, "Native App (Snowflake sign-in)", size=11, color=RGBColor.from_string(color), after=2)
        link_table(doc, ["Region", "App URL", "Version"], rws, widths=[1.4, 4.2, 1.3])

    if lst:
        rws = []
        for n in lst:
            for conn, lbl in REG:
                r = listing(conn, n)
                rws.append([("Native App" if r and r["is_application"] == "true" else "Data share") if conn == REG[0][0] else "",
                            lbl, r["global_name"] if r else "not published", r["uniform_listing_locator"] if r else "—"])
        body(doc, "Internal Marketplace listings", size=11, color=RGBColor.from_string(color), after=2)
        table(doc, ["Listing", "Region", "Global name", "Uniform listing locator"], rws, widths=[1.0, 1.4, 1.3, 3.2], size=8.5)

    files = kit_files(kit)
    if files:
        body(doc, "Presales kit", size=11, color=RGBColor.from_string(color), after=2)
        body(doc, str(SAP / kit), size=8.5, color=GREY, after=2)
        for f in files:
            bullet(doc, f, size=9)

# 5. shared
page_break(doc)
h1(doc, "5. Shared tooling, skills and cross-demo documents")
link_table(doc, ["Item", "Location", "Purpose"], [
    ["Shared Word / PowerPoint kits", "sap-bdc-finance-360/tools/sap_docx_kit.py, sap_pptx_kit.py", "Branding and slide checks used by every kit"],
    ["Screenshot tool", "sap-bdc-finance-360/tools/capture_shots.py", "Captures every app page for decks"],
    ["Portfolio capture", "sap-bdc-spend-360/tools/capture_portfolio.py", "Screenshots and KPIs for this index"],
    ["Narrated video engine", "sap-bdc-finance-360/tools/video/build.py + segments_<domain>.py", "Builds the walkthrough MP4s"],
    ["SAP BDC data products", ("url", GH + "sap-bdc-data-products", GH + "sap-bdc-data-products"), "Data product catalog and CSN"],
    ["SAP BDC CDC", ("url", GH + "sap-bdc-cdc", GH + "sap-bdc-cdc"), "CDC recommendation for BDC shares"],
    ["SAP BDC MCP server", ("url", GH + "sap-bdc-snowflake-mcp-server", GH + "sap-bdc-snowflake-mcp-server"), "MCP tools for BDC Connect on Snowflake"],
    ["Use case catalog", str(HOME / "Downloads/SAP_BDC_360_Use_Case_Catalog.docx"), "Use cases across the 360 apps"],
    ["Marketplace listing index", str(HOME / "Downloads/SAP_BDC_360_Marketplace_Listings.docx"), "Listings in all three regions"],
    ["Skill: sap-bdc-360-app-lifecycle", "~/.snowflake/cortex/skills/sap-bdc-360-app-lifecycle", "Build, ship, publish and kit a 360 app"],
    ["Skill: launch_core_apps", "~/.snowflake/cortex/skills/launch_core_apps", "Start every local app on its port"],
    ["Skill: provision-360-access", "~/.snowflake/cortex/skills/provision-360-access", "Give a user access to a 360 Native App"],
], widths=[1.8, 3.3, 1.8])

# 6. gaps
h1(doc, "6. Known gaps")
gaps = []
for k, title, _, repo, pub, local, ports, appn, lst, kit in P:
    if appn:
        missing = [lbl for conn, lbl in REG if not app(conn, appn)]
        if missing:
            gaps.append(f"{title}: Native App not installed in {', '.join(missing)}.")
        unv = [lbl for conn, lbl in REG if (app(conn, appn) or {}).get("version") == "UNVERSIONED"]
        if unv:
            gaps.append(f"{title}: Native App is installed unversioned (development mode) in {', '.join(unv)}; upgrades go through a versioned patch.")
    if not repo:
        gaps.append(f"{title}: no source repository, public site, Native App or presales kit yet; local only.")
    elif VIS.get(repo) == "Private":
        gaps.append(f"{title}: source repository {repo} is private.")
    if kit and not any(f.endswith(".mp4") for f in kit_files(kit)):
        gaps.append(f"{title}: presales kit has no narrated walkthrough video.")
    if not PF.get(k):
        gaps.append(f"{title}: dashboard could not be captured for this index.")
for g in gaps:
    bullet(doc, g)

doc.save(OUT)
print(OUT, f"{OUT.stat().st_size / 1e6:.1f} MB", len(gaps), "gaps")
